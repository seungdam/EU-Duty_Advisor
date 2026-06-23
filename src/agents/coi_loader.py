"""Read COI spreadsheets as compact classifier evidence.

The COI files are source evidence, not generated workbook artifacts.  This
loader intentionally flattens visible cell values into bounded plain text so
classification smoke tests can compare OCR-only vs OCR+COI inputs without
depending on a full COI schema yet.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_COI_ROOT = PROJECT_ROOT / "test" / "COI(식품원재료풀이)"

TOKEN_RE = re.compile(r"[0-9A-Za-z가-힣]+")
INDEX_PREFIX_RE = re.compile(r"^\s*([0-9,\s]+)\.")


@dataclass(frozen=True)
class CoiEvidence:
    path: Path
    text: str
    matched_score: int = 0


def normalize_text(value: object) -> str:
    return unicodedata.normalize("NFC", " ".join(str(value or "").split()))


def evidence_tokens(text: str) -> set[str]:
    text = normalize_text(text)
    return {
        token.lower()
        for token in TOKEN_RE.findall(text or "")
        if len(token) >= 2
    }


def _indices_from_parent(path: Path) -> set[int]:
    match = INDEX_PREFIX_RE.match(path.parent.name)
    if not match:
        return set()
    out: set[int] = set()
    for chunk in match.group(1).replace(" ", "").split(","):
        if chunk.isdigit():
            out.add(int(chunk))
    return out


def _xlsx_paths(root: Path) -> list[Path]:
    if not root.exists():
        return []
    return sorted(
        path
        for path in root.rglob("*.xlsx")
        if not path.name.startswith("~$")
    )


def _score_path(path: Path, *, case_index: int | None, product_name: str) -> int:
    score = 0
    if case_index is not None and case_index in _indices_from_parent(path):
        score += 100
    haystack = f"{path.parent.name} {path.stem}"
    overlap = evidence_tokens(product_name) & evidence_tokens(haystack)
    score += len(overlap) * 10

    # Prefer final/exportable files over intermediate revision copies when tied.
    lowered = path.name.lower()
    if "수출가능" in path.name:
        score += 4
    if "수정" in path.name or "최종" in path.name:
        score += 2
    if "복사" in path.name or "copy" in lowered:
        score -= 1
    return score


def find_coi_path(
    *,
    case_index: int | None,
    product_name: str,
    coi_root: Path = DEFAULT_COI_ROOT,
) -> Path | None:
    candidates = _xlsx_paths(coi_root)
    if not candidates:
        return None
    scored = [
        (_score_path(path, case_index=case_index, product_name=product_name), path)
        for path in candidates
    ]
    scored.sort(key=lambda item: (-item[0], str(item[1])))
    best_score, best_path = scored[0]
    return best_path if best_score > 0 else None


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    text = normalize_text(value)
    if not text:
        return ""
    if text.lower() in {"none", "nan"}:
        return ""
    return text


def flatten_xlsx_text(
    path: Path,
    *,
    max_chars: int = 8000,
    max_rows_per_sheet: int = 220,
) -> str:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required to read COI xlsx files") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows_added = 0
            parts.append(f"[sheet] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_to_text(value) for value in row]
                values = [value for value in values if value]
                if not values:
                    continue
                parts.append(" | ".join(values))
                rows_added += 1
                if rows_added >= max_rows_per_sheet:
                    parts.append("[sheet_truncated]")
                    break
                if sum(len(part) + 1 for part in parts) >= max_chars:
                    return "\n".join(parts)[:max_chars]
    finally:
        workbook.close()
    return "\n".join(parts)[:max_chars]


@lru_cache(maxsize=256)
def _cached_flatten(path_text: str, max_chars: int) -> str:
    return flatten_xlsx_text(Path(path_text), max_chars=max_chars)


def load_coi_evidence(
    *,
    case_index: int | None,
    product_name: str,
    coi_root: Path = DEFAULT_COI_ROOT,
    max_chars: int = 8000,
) -> CoiEvidence | None:
    path = find_coi_path(
        case_index=case_index,
        product_name=product_name,
        coi_root=coi_root,
    )
    if path is None:
        return None
    text = _cached_flatten(str(path), max_chars)
    score = _score_path(path, case_index=case_index, product_name=product_name)
    return CoiEvidence(path=path, text=text, matched_score=score)


def summarize_coi_matches(
    rows: Iterable[tuple[int | None, str]],
    *,
    coi_root: Path = DEFAULT_COI_ROOT,
) -> list[dict[str, object]]:
    out: list[dict[str, object]] = []
    for case_index, product_name in rows:
        path = find_coi_path(
            case_index=case_index,
            product_name=product_name,
            coi_root=coi_root,
        )
        out.append(
            {
                "case_index": case_index,
                "product_name": product_name,
                "coi_path": str(path) if path else "",
            }
        )
    return out
